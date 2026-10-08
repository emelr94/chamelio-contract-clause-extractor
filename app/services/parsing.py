"""Turn an uploaded file into per-page plain text. No LLM involved."""

import io
from dataclasses import dataclass

import docx
import pymupdf

from app.errors import AppError
from app.schemas.enums import ErrorCode, FileType

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"  # DOCX is a zip container
EXTENSIONS: dict[str, FileType] = {".pdf": FileType.PDF, ".docx": FileType.DOCX}
NBSP = "\u00a0"


@dataclass(frozen=True)
class ParsedDocument:
    file_type: FileType
    pages: list[str]
    has_page_numbers: bool  # False for DOCX: Word has no fixed pagination

    @property
    def page_count(self) -> int | None:
        return len(self.pages) if self.has_page_numbers else None


def detect_file_type(filename: str, content: bytes) -> FileType:
    """Require the extension and the magic bytes to agree, so a renamed file is rejected."""
    extension = filename[filename.rfind(".") :].lower() if "." in filename else ""
    file_type = EXTENSIONS.get(extension)
    magic = PDF_MAGIC if file_type is FileType.PDF else ZIP_MAGIC
    if file_type is None or not content.startswith(magic):
        raise AppError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            "Only PDF (.pdf) and Word (.docx) files are supported.",
            {"filename": filename},
        )
    return file_type


def parse_document(
    file_type: FileType, content: bytes, *, max_pages: int, min_chars_per_page: int
) -> ParsedDocument:
    if file_type is FileType.PDF:
        return _parse_pdf(content, max_pages=max_pages, min_chars_per_page=min_chars_per_page)
    return _parse_docx(content)


def _parse_pdf(content: bytes, *, max_pages: int, min_chars_per_page: int) -> ParsedDocument:
    try:
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise AppError(ErrorCode.UNREADABLE_DOCUMENT, "The PDF is password-protected.")
            if pdf.page_count == 0:
                raise AppError(ErrorCode.UNREADABLE_DOCUMENT, "The PDF has no pages.")
            if pdf.page_count > max_pages:
                raise AppError(
                    ErrorCode.DOCUMENT_TOO_LARGE,
                    f"The PDF has {pdf.page_count} pages; the limit is {max_pages}.",
                    {"page_count": pdf.page_count, "max_pages": max_pages},
                )
            pages = [page.get_text().replace(NBSP, " ") for page in pdf]
    except AppError:
        raise
    except Exception as exc:  # PyMuPDF raises several unrelated types for corrupt input
        raise AppError(ErrorCode.UNREADABLE_DOCUMENT, "The PDF could not be opened.") from exc

    total_chars = sum(len(page.strip()) for page in pages)
    if total_chars < min_chars_per_page * len(pages):
        raise AppError(
            ErrorCode.SCANNED_PDF_UNSUPPORTED,
            "The PDF has little or no extractable text; it is probably scanned. "
            "OCR is not supported yet.",
            {"page_count": len(pages), "extracted_chars": total_chars},
        )
    return ParsedDocument(FileType.PDF, pages, has_page_numbers=True)


def _parse_docx(content: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise AppError(ErrorCode.UNREADABLE_DOCUMENT, "The DOCX could not be opened.") from exc
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip()).replace(NBSP, " ")
    if not text:
        raise AppError(ErrorCode.UNREADABLE_DOCUMENT, "The DOCX contains no text.")
    return ParsedDocument(FileType.DOCX, [text], has_page_numbers=False)
